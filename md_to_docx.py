# -*- coding: utf-8 -*-
"""Minimal, dependency-light Markdown -> DOCX converter tuned for this manuscript.

Handles: # / ## / ### headings, paragraphs, bullet and numbered lists,
markdown tables, and inline **bold** / *italic*. Produces a journal-style
manuscript (Times New Roman 12 pt, double spaced, page numbers, line numbers).
"""
import io, os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml

BOLD = re.compile(r'\*\*(.+?)\*\*')
ITAL = re.compile(r'(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)')


def set_style_font(style, name, size=None, bold=None):
    style.font.name = name
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rFonts.set(qn(a), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def clean_heading(style):
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rFonts.get(qn(attr)):
                del rFonts.attrib[qn(attr)]
    pPr = style.element.find(qn('w:pPr'))
    if pPr is not None:
        pBdr = pPr.find(qn('w:pBdr'))
        if pBdr is not None:
            pPr.remove(pBdr)


def add_runs(par, text):
    """Render inline markdown (bold / italic) into paragraph runs."""
    tokens = []
    pos = 0
    pattern = re.compile(r'\*\*(.+?)\*\*|(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)')
    for m in pattern.finditer(text):
        if m.start() > pos:
            tokens.append((text[pos:m.start()], False, False))
        if m.group(1) is not None:
            tokens.append((m.group(1), True, False))
        else:
            tokens.append((m.group(2), False, True))
        pos = m.end()
    if pos < len(text):
        tokens.append((text[pos:], False, False))
    if not tokens:
        tokens = [(text, False, False)]
    for t, b, i in tokens:
        if t == '':
            continue
        r = par.add_run(t)
        r.bold = b
        r.italic = i


def is_table_row(line):
    return line.strip().startswith('|') and line.strip().endswith('|')


def is_sep_row(line):
    return bool(re.fullmatch(r'\|[\s:\-|]+\|', line.strip()))


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def build(md_path, out_path, title_center=True, double=True):
    text = io.open(md_path, encoding='utf-8').read()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.54)

    normal = doc.styles['Normal']
    set_style_font(normal, 'Times New Roman', 12)
    normal.paragraph_format.line_spacing = 2.0 if double else 1.15
    normal.paragraph_format.space_after = Pt(0)

    clean_heading(doc.styles['Title'])
    set_style_font(doc.styles['Title'], 'Times New Roman', 15, True)
    for i in (1, 2, 3):
        st = doc.styles['Heading %d' % i]
        clean_heading(st)
        set_style_font(st, 'Times New Roman', 13 - i, True)
    for st in ('List Bullet', 'List Number'):
        set_style_font(doc.styles[st], 'Times New Roman', 12)

    # page numbers, centred footer
    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    r._r.append(parse_xml('<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w')))
    r = fp.add_run()
    r._r.append(parse_xml('<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w')))
    r = fp.add_run()
    r._r.append(parse_xml('<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w')))

    # continuous line numbering (Elsevier reviewers like it)
    try:
        sectPr = sec._sectPr
        ln = parse_xml('<w:lnNumType %s w:countBy="1" w:restart="continuous"/>' % nsdecls('w'))
        sectPr.append(ln)
    except Exception as e:
        print('  (line numbering skipped:', e, ')')

    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == '---':
            i += 1
            continue
        # table block
        if is_table_row(lines[i]):
            block = []
            while i < len(lines) and is_table_row(lines[i]):
                block.append(lines[i])
                i += 1
            rows = [split_row(b) for b in block if not is_sep_row(b)]
            if rows:
                ncol = max(len(r_) for r_ in rows)
                t = doc.add_table(rows=0, cols=ncol)
                t.style = 'Table Grid'
                for k, cells in enumerate(rows):
                    cells = cells + [''] * (ncol - len(cells))
                    row = t.add_row()
                    for c, txt in zip(row.cells, cells):
                        p = c.paragraphs[0]
                        p.paragraph_format.line_spacing = 1.0
                        p.paragraph_format.space_after = Pt(0)
                        add_runs(p, txt)
                        if k == 0:
                            for run in p.runs:
                                run.bold = True
            doc.add_paragraph()
            continue
        # headings
        m = re.match(r'^(#{1,3})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1))
            if lvl == 1:
                h = doc.add_heading(level=0)
            else:
                h = doc.add_heading(level=lvl - 1)
            add_runs(h, m.group(2))
            if lvl == 1 and title_center:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        # lists
        m = re.match(r'^[-*]\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            i += 1
            continue
        # plain paragraph
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        add_runs(p, s)
        i += 1

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print('wrote', out_path, os.path.getsize(out_path), 'bytes')


if __name__ == '__main__':
    BASE = '/sandbox/workspace/heavymetal'
    OUT = os.path.join(BASE, 'submission')
    build(os.path.join(BASE, 'manuscript.md'), os.path.join(OUT, 'manuscript.docx'))
    build(os.path.join(BASE, 'highlights.md'), os.path.join(OUT, 'highlights.docx'), title_center=False, double=False)
    build(os.path.join(BASE, 'cover_letter_EnvironmentalResearch.md'), os.path.join(OUT, 'cover_letter.docx'), title_center=False, double=False)
